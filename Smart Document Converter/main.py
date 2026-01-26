from fpdf import FPDF
from docx import Document

def txt_to_pdf(txt_file):
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Arial", size=12)

    with open(txt_file, 'r') as f:
        for line in f:
            pdf.cell(200, 10, txt=line.encode('latin-1', 'ignore').decode('latin-1'), ln=True)

    output = txt_file.replace('.txt', '.pdf')
    pdf.output(output)
    print(f"✅ Converted to {output}")

def docx_to_txt(docx_file):
    doc = Document(docx_file)
    output = docx_file.replace('.docx', '.txt')
    with open(output, 'w') as f:
        for para in doc.paragraphs:
            f.write(para.text + '\n')
    print(f"✅ Converted to {output}")

def main():
    print("📑 Document Converter")
    print("1. TXT → PDF")
    print("2. DOCX → TXT")
    choice = input("Choose (1/2): ")

    file = input("Enter file path: ")

    if choice == '1':
        txt_to_pdf(file)
    elif choice == '2':
        docx_to_txt(file)
    else:
        print("❌ Invalid choice!")

if __name__ == "__main__":
    main()